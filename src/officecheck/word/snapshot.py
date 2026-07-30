"""Word (.docx) snapshot extraction (read-only)."""

from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

from docx import Document


def _section(sec) -> dict:
    def emu(v):
        return int(v) if v is not None else None

    return {
        "orientation": str(sec.orientation),
        "page_width": emu(sec.page_width),
        "page_height": emu(sec.page_height),
        "margins": (emu(sec.left_margin), emu(sec.right_margin),
                    emu(sec.top_margin), emu(sec.bottom_margin)),
        "header_distance": emu(sec.header_distance),
        "footer_distance": emu(sec.footer_distance),
    }


def _header_footer_text(sec) -> dict:
    out = {}
    for name, part in (("header", sec.header), ("footer", sec.footer)):
        if part is None or part.is_linked_to_previous:
            continue
        text = "\n".join(p.text for p in part.paragraphs).strip()
        if text:
            out[name] = text
    return out


def take_snapshot(path: str | Path) -> dict:
    path = Path(path)
    doc = Document(str(path))

    paragraphs = [
        {"text": p.text, "style": p.style.name if p.style else None}
        for p in doc.paragraphs
    ]

    tables = []
    for table in doc.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        tables.append({
            "style": table.style.name if table.style else None,
            "rows": rows,
        })

    sections = [_section(sec) for sec in doc.sections]
    header_footer = [_header_footer_text(sec) for sec in doc.sections]

    images = []
    with zipfile.ZipFile(path) as zf:
        for name in zf.namelist():
            if name.startswith("word/media/"):
                images.append(hashlib.sha256(zf.read(name)).hexdigest())

    return {
        "target": "word",
        "format": "docx",
        "paragraphs": paragraphs,
        "tables": tables,
        "sections": sections,
        "header_footer": header_footer,
        "images": sorted(images),
    }
